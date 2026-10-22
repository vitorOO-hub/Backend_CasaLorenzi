"""Le e prepara a carga historica de estoque Vulto sem alterar a planilha."""

import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from .catalogo_provisorio import planejar_placeholders_catalogo
from .lojas import CODIGO_LOJA_POR_ID_VULTO, resolver_id_loja_vulto

ORIGEM_MOVIMENTOS = "vulto_entradas_conferencia"
CODIGO_LOTE = "VULTO-ESTOQUE-20261022"
COLECAO_MIGRACAO = "migracao vulto"
LOJA_POR_ORIGEM_TRANSFERENCIA = {
    "pinheiros": 1,
    "eldorado": 2,
    "savassi": 3,
    "batel": 4,
    "ecommerce": 5,
    "app": 6,
    "cd cajamar": 7,
}


class ImportacaoEstoqueError(RuntimeError):
    pass


def _texto(valor: object) -> str:
    return "" if valor is None else str(valor).strip()


def _inteiro(valor: object, campo: str) -> int:
    try:
        decimal = Decimal(_texto(valor))
    except (InvalidOperation, ValueError) as erro:
        raise ImportacaoEstoqueError(f"{campo} invalido: {valor!r}") from erro
    if decimal != decimal.to_integral_value():
        raise ImportacaoEstoqueError(f"{campo} precisa ser inteiro: {valor!r}")
    return int(decimal)


def _dinheiro(valor: object, campo: str) -> Decimal | None:
    texto = _texto(valor)
    if not texto:
        return None
    try:
        if "," in texto:
            texto = texto.replace(".", "").replace(",", ".")
        numero = Decimal(texto).quantize(Decimal("0.01"))
        if not numero.is_finite():
            raise InvalidOperation
        return numero
    except InvalidOperation as erro:
        raise ImportacaoEstoqueError(f"{campo} invalido: {valor!r}") from erro


def _data_hora(valor: object, campo: str) -> datetime:
    if isinstance(valor, datetime):
        resultado = valor
    else:
        texto = _texto(valor).replace("Z", "+00:00")
        try:
            resultado = datetime.fromisoformat(texto)
        except ValueError as erro:
            raise ImportacaoEstoqueError(f"{campo} ausente ou invalido: {valor!r}") from erro
    if resultado.tzinfo is None:
        raise ImportacaoEstoqueError(f"{campo} precisa ter fuso horario: {valor!r}")
    return resultado


def _linhas(aba: Any) -> list[dict[str, object]]:
    linhas = aba.iter_rows(values_only=True)
    cabecalho = next(linhas, None)
    if not cabecalho:
        raise ImportacaoEstoqueError(f"A aba {aba.title} esta vazia")
    nomes = [_texto(campo) for campo in cabecalho]
    if not all(nomes) or len(set(nomes)) != len(nomes):
        raise ImportacaoEstoqueError(f"Cabecalho invalido na aba {aba.title}")
    resultado = []
    for numero, valores in enumerate(linhas, start=2):
        if not any(valor is not None for valor in valores):
            continue
        if len(valores) != len(nomes):
            raise ImportacaoEstoqueError(f"Colunas inesperadas na aba {aba.title}, linha {numero}")
        resultado.append(dict(zip(nomes, valores, strict=True)))
    return resultado


def carregar_planilha(caminho: Path) -> dict[str, list[dict[str, object]]]:
    """Carrega somente as abas finais usadas nesta importacao."""
    try:
        from openpyxl import load_workbook
    except ImportError as erro:
        raise ImportacaoEstoqueError("Instale openpyxl para ler a planilha") from erro

    workbook = load_workbook(caminho, read_only=True, data_only=True)
    obrigatorias = {
        "estoque_conferencia",
        "entradas_conferencia",
        "catalogo_adicoes_final",
        "itens_sem_sku_real",
    }
    faltantes = obrigatorias - set(workbook.sheetnames)
    if faltantes:
        raise ImportacaoEstoqueError(
            "Abas obrigatorias ausentes: " + ", ".join(sorted(faltantes))
        )
    return {nome: _linhas(workbook[nome]) for nome in obrigatorias}


def _sku_linha(linha: dict[str, object], campo: str = "sku_padronizado") -> str:
    sku = _texto(linha.get(campo))
    if not sku:
        raise ImportacaoEstoqueError(f"Linha sem {campo}: {linha.get('id', '')}")
    return sku


def _id_origem_movimento(linha: dict[str, object]) -> str:
    id_origem = _texto(linha.get("id"))
    if id_origem:
        return id_origem
    observacao = _texto(linha.get("obs"))
    envio = re.search(r"deduzida do envio id\s+(\d+)", observacao, flags=re.IGNORECASE)
    if not envio:
        raise ImportacaoEstoqueError(
            "Movimento sem ID legado e sem referencia ao envio original: "
            f"{observacao or linha!r}"
        )
    return f"derivado-envio-{envio.group(1)}"


def preparar_snapshot(
    linhas: list[dict[str, object]], lojas_por_codigo: dict[str, object]
) -> list[dict[str, object]]:
    snapshot = []
    chaves = set()
    for linha in linhas:
        loja_legada = _inteiro(linha.get("loja_id"), "estoque_conferencia.loja_id")
        try:
            id_loja = resolver_id_loja_vulto(loja_legada, lojas_por_codigo)
        except ValueError as erro:
            raise ImportacaoEstoqueError(str(erro)) from erro
        sku = _sku_linha(linha)
        quantidade = _inteiro(linha.get("qtd"), f"estoque {sku}.qtd")
        atualizado = _data_hora(
            linha.get("data_padronizada") or linha.get("atualizado_em"),
            f"estoque {sku}.data_padronizada",
        )
        chave = (loja_legada, sku)
        if chave in chaves:
            raise ImportacaoEstoqueError(f"Snapshot duplicado para loja {loja_legada}, SKU {sku}")
        chaves.add(chave)
        snapshot.append(
            {
                "loja_id_vulto": loja_legada,
                "id_loja": id_loja,
                "sku": sku,
                "quantidade": quantidade,
                "atualizado_em": atualizado,
                "saldo_inferido_por_ausencia": False,
            }
        )
    return snapshot


def completar_snapshot_com_saldos_zero(
    snapshot: list[dict[str, object]],
    linhas_movimentos: list[dict[str, object]],
    lojas_por_codigo: dict[str, object],
) -> list[dict[str, object]]:
    """Materializa como zero as chaves ausentes, conforme decisao aprovada."""
    resultado = list(snapshot)
    chaves_existentes = {
        (linha["loja_id_vulto"], linha["sku"]) for linha in resultado
    }
    datas_por_chave: dict[tuple[int, str], datetime] = {}
    lojas_por_chave: dict[tuple[int, str], object] = {}
    for linha in linhas_movimentos:
        loja_legada = _inteiro(linha.get("loja_id"), "entradas_conferencia.loja_id")
        sku = _sku_linha(linha)
        chave = (loja_legada, sku)
        if chave in chaves_existentes:
            continue
        data_movimento = _data_hora(
            linha.get("data_padronizada"),
            f"movimento {_texto(linha.get('id'))}.data_padronizada",
        )
        datas_por_chave[chave] = max(
            datas_por_chave.get(chave, data_movimento), data_movimento
        )
        try:
            lojas_por_chave[chave] = resolver_id_loja_vulto(
                loja_legada, lojas_por_codigo
            )
        except ValueError as erro:
            raise ImportacaoEstoqueError(str(erro)) from erro

    for (loja_legada, sku), atualizado in sorted(datas_por_chave.items()):
        resultado.append(
            {
                "loja_id_vulto": loja_legada,
                "id_loja": lojas_por_chave[(loja_legada, sku)],
                "sku": sku,
                "quantidade": 0,
                "atualizado_em": atualizado,
                "saldo_inferido_por_ausencia": True,
            }
        )
    return resultado


def preparar_catalogo(
    adicoes: list[dict[str, object]],
    itens_temporarios: list[dict[str, object]],
    skus_necessarios: set[str],
    skus_existentes: set[str],
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    """Planeja somente SKUs ausentes; nunca atualiza produto/variacao existente."""
    faltantes = skus_necessarios - skus_existentes
    adicoes_por_sku = {}
    for linha in adicoes:
        sku = _texto(linha.get("sku"))
        if not sku or sku in adicoes_por_sku:
            raise ImportacaoEstoqueError(f"SKU vazio ou duplicado na aba de adicoes: {sku!r}")
        adicoes_por_sku[sku] = linha
    temporarios = {
        _texto(linha.get("sku_provisorio"))
        for linha in itens_temporarios
        if _texto(linha.get("sku_provisorio"))
    }
    sem_origem = faltantes - set(adicoes_por_sku) - temporarios
    if sem_origem:
        raise ImportacaoEstoqueError(
            "SKUs ausentes sem cadastro aprovado ou provisao: "
            + ", ".join(sorted(sem_origem)[:20])
        )

    produtos_por_chave: dict[tuple[str, str], dict[str, object]] = {}
    variacoes: list[dict[str, object]] = []
    for sku in sorted(faltantes & set(adicoes_por_sku)):
        linha = adicoes_por_sku[sku]
        nome = _texto(linha.get("produto"))
        categoria = _texto(linha.get("categoria"))
        cor = _texto(linha.get("cor"))
        tamanho = _texto(linha.get("tamanho"))
        preco = _dinheiro(linha.get("preco"), f"{sku}.preco")
        custo = _dinheiro(linha.get("custo"), f"{sku}.custo")
        acao = _texto(linha.get("acao_importacao"))
        if not all((nome, categoria, cor, tamanho)) or preco is None:
            raise ImportacaoEstoqueError(f"Cadastro aprovado incompleto para {sku}")
        if acao not in {"cadastrar", "cadastrar_como_inativo"}:
            raise ImportacaoEstoqueError(f"Acao de catalogo nao reconhecida para {sku}: {acao}")
        chave_produto = ("Vulto", nome)
        ativo = acao == "cadastrar"
        produto = produtos_por_chave.get(chave_produto)
        if produto is None:
            produto = {
                "marca": "Vulto",
                "nome": nome,
                "categoria": categoria,
                "colecao": COLECAO_MIGRACAO,
                "preco_base": preco,
                "ativo": ativo,
                "cores": {cor},
            }
            produtos_por_chave[chave_produto] = produto
        elif produto["categoria"] != categoria:
            raise ImportacaoEstoqueError(
                f"Produto {nome!r} tem mais de uma categoria na aba de adicoes"
            )
        else:
            produto["preco_base"] = min(produto["preco_base"], preco)
            produto["ativo"] = bool(produto["ativo"] or ativo)
            produto["cores"].add(cor)
        if any(
            atual["produto_chave"] == chave_produto
            and atual["cor"] == cor
            and atual["tamanho"] == tamanho
            for atual in variacoes
        ):
            raise ImportacaoEstoqueError(
                f"Grade duplicada no catalogo aprovado para {sku}: {cor}/{tamanho}"
            )
        variacoes.append(
            {
                "sku": sku,
                "cor": cor,
                "tamanho": tamanho,
                "preco_venda": preco,
                "custo": custo,
                "ean": _texto(linha.get("ean")) or None,
                "ativa": ativo,
                "produto_chave": chave_produto,
            }
        )

    placeholders = planejar_placeholders_catalogo(
        temporarios, [{"sku": sku} for sku in skus_existentes]
    )
    grades_planejadas = {
        (variacao["produto_chave"], variacao["cor"], variacao["tamanho"]): variacao["sku"]
        for variacao in variacoes
    }
    for item in placeholders:
        if item["sku"] not in faltantes:
            continue
        chave_produto = (str(item["marca"]), str(item["nome_produto"]))
        chave_grade = (chave_produto, item["cor"], item["tamanho"])
        if chave_grade in grades_planejadas:
            raise ImportacaoEstoqueError(
                f"SKU provisorio {item['sku']} colide com a grade "
                f"{grades_planejadas[chave_grade]} ({item['cor']}/{item['tamanho']})"
            )
        grades_planejadas[chave_grade] = item["sku"]
        produtos_por_chave[chave_produto] = {
            "marca": item["marca"],
            "nome": item["nome_produto"],
            "categoria": item["categoria"],
            "colecao": item["colecao"],
            "preco_base": item["preco_base"],
            "ativo": item["produto_ativo"],
            "cores": {item["cor"]},
        }
        variacoes.append(
            {
                "sku": item["sku"],
                "cor": item["cor"],
                "tamanho": item["tamanho"],
                "preco_venda": item["preco_venda"],
                "custo": None,
                "ean": None,
                "ativa": item["variacao_ativa"],
                "produto_chave": chave_produto,
            }
        )

    produtos = list(produtos_por_chave.values())
    for produto in produtos:
        produto["cores"] = sorted(produto["cores"])
    return produtos, variacoes


def preparar_movimentos(
    linhas: list[dict[str, object]],
    snapshot: list[dict[str, object]],
    lojas_por_codigo: dict[str, object],
) -> list[dict[str, object]]:
    """Reconstrui saldos encadeados; bloqueia chaves sem saldo final oficial."""
    saldos_finais = {
        (linha["loja_id_vulto"], linha["sku"]): linha["quantidade"] for linha in snapshot
    }
    normalizados = []
    ids_vistos = set()
    por_chave: dict[tuple[int, str], list[dict[str, object]]] = defaultdict(list)

    for linha in linhas:
        id_origem = _id_origem_movimento(linha)
        if not id_origem or id_origem in ids_vistos:
            raise ImportacaoEstoqueError(f"ID de movimento vazio ou duplicado: {id_origem!r}")
        ids_vistos.add(id_origem)
        sku = _sku_linha(linha)
        loja_legada = _inteiro(linha.get("loja_id"), "entradas_conferencia.loja_id")
        try:
            id_loja = resolver_id_loja_vulto(loja_legada, lojas_por_codigo)
        except ValueError as erro:
            raise ImportacaoEstoqueError(str(erro)) from erro
        quantidade = _inteiro(linha.get("qtd"), f"movimento {id_origem}.qtd")
        if quantidade <= 0:
            raise ImportacaoEstoqueError(f"Movimento {id_origem} precisa ter quantidade positiva")
        codigo_origem = _texto(linha.get("tipo_movimento"))
        obs = _texto(linha.get("obs"))
        if codigo_origem == "saida_transferencia":
            tipo = "transferencia_saida"
            sinal = -1
        elif codigo_origem == "entrada":
            eh_transferencia = obs.casefold().startswith("veio da ")
            tipo = "transferencia_entrada" if eh_transferencia else "entrada"
            sinal = 1
        else:
            raise ImportacaoEstoqueError(
                f"Tipo de movimento nao reconhecido na linha {id_origem}: {codigo_origem}"
            )
        destino = _texto(linha.get("loja_destino"))
        id_loja_destino = None
        if codigo_origem == "saida_transferencia" and not destino:
            raise ImportacaoEstoqueError(
                f"Saida de transferencia {id_origem} sem loja_destino"
            )
        if codigo_origem == "entrada" and obs.casefold().startswith("veio da "):
            nome_origem = obs.casefold().removeprefix("veio da ").strip()
            loja_origem = LOJA_POR_ORIGEM_TRANSFERENCIA.get(nome_origem)
            if loja_origem is None:
                raise ImportacaoEstoqueError(
                    f"Origem da transferencia nao reconhecida no movimento {id_origem}: {obs}"
                )
            resolver_id_loja_vulto(loja_origem, lojas_por_codigo)
        if destino:
            try:
                id_loja_destino = resolver_id_loja_vulto(
                    _inteiro(destino, "loja_destino"), lojas_por_codigo
                )
            except ValueError as erro:
                raise ImportacaoEstoqueError(str(erro)) from erro
        if tipo == "transferencia_entrada":
            id_loja_destino = None
        item = {
            "id_origem": id_origem,
            "id_loja": id_loja,
            "id_loja_destino": id_loja_destino,
            "loja_id_vulto": loja_legada,
            "sku": sku,
            "codigo_tipo": tipo,
            "sinal": sinal,
            "quantidade": quantidade,
            "criada_em": _data_hora(
                linha.get("data_padronizada"), f"movimento {id_origem}.data_padronizada"
            ),
            "motivo": obs or f"Movimento legado Vulto {tipo}",
            "custo_unitario": _dinheiro(
                linha.get("custo_unit"), f"movimento {id_origem}.custo_unit"
            ),
            "nota_fiscal": _texto(linha.get("nf")) or None,
            "fornecedor": _texto(linha.get("fornecedor")) or None,
            "dados_origem": dict(linha),
        }
        normalizados.append(item)
        por_chave[(loja_legada, sku)].append(item)

    chaves_sem_snapshot = sorted(set(por_chave) - set(saldos_finais))
    if chaves_sem_snapshot:
        raise ImportacaoEstoqueError(
            f"{len(chaves_sem_snapshot)} combinacoes sem saldo final; "
            "complete o snapshot com uma decisao aprovada antes de importar"
        )
    chaves_inferidas = {
        (linha["loja_id_vulto"], linha["sku"])
        for linha in snapshot
        if linha.get("saldo_inferido_por_ausencia")
    }

    for chave, movimentos in por_chave.items():
        saldo_final = saldos_finais[chave]
        saldo_abertura = saldo_final - sum(m["sinal"] * m["quantidade"] for m in movimentos)
        saldo_atual = saldo_abertura
        for movimento in sorted(
            movimentos,
            key=lambda m: (
                m["criada_em"],
                int(m["id_origem"].removeprefix("derivado-envio-")),
            ),
        ):
            movimento["quantidade_anterior"] = saldo_atual
            saldo_atual += movimento["sinal"] * movimento["quantidade"]
            movimento["quantidade_posterior"] = saldo_atual
            movimento["saldo_reconstruido"] = True
            movimento["snapshot_ausente_tratado_como_zero"] = chave in chaves_inferidas
        if saldo_atual != saldo_final:
            raise ImportacaoEstoqueError(f"Saldo reconstruido nao fechou para {chave}")

    return sorted(
        normalizados,
        key=lambda m: (
            m["criada_em"],
            int(m["id_origem"].removeprefix("derivado-envio-")),
        ),
    )


def _json_seguro(valor: object) -> object:
    if isinstance(valor, (datetime, date)):
        return valor.isoformat()
    if isinstance(valor, Decimal):
        return str(valor)
    if valor is None or isinstance(valor, (str, int, float, bool)):
        return valor
    return str(valor)


def _produto_variacao_no_banco(
    cursor: Any, produto: dict[str, object], id_unidade: object
) -> object:
    cursor.execute(
        """
        SELECT id_produto, categoria FROM public.produto
        WHERE nome = %s AND marca = %s FOR UPDATE
        """,
        (produto["nome"], produto["marca"]),
    )
    existente = cursor.fetchone()
    if existente:
        categoria_existente = _texto(existente[1])
        if categoria_existente and categoria_existente != produto["categoria"]:
            raise ImportacaoEstoqueError(
                f"Categoria divergente no produto existente {produto['nome']!r}: "
                f"banco={categoria_existente!r}, planilha={produto['categoria']!r}"
            )
        return existente[0]

    from psycopg.types.json import Jsonb

    cursor.execute(
        """
        INSERT INTO public.produto (
            nome, marca, categoria, colecao, preco_base, ativo, cores, id_unidade
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
        RETURNING id_produto
        """,
        (
            produto["nome"],
            produto["marca"],
            produto["categoria"],
            produto["colecao"],
            produto["preco_base"],
            produto["ativo"],
            Jsonb(produto["cores"]),
            id_unidade,
        ),
    )
    return cursor.fetchone()[0]


def gravar_importacao(conexao: Any, fontes: dict[str, list[dict[str, object]]]) -> dict[str, int]:
    """Grava catalogo ausente, snapshot e movimentos Vulto numa unica transacao."""
    from psycopg.types.json import Jsonb

    with conexao.transaction():
        with conexao.cursor() as cursor:
            cursor.execute("SELECT pg_advisory_xact_lock(hashtext('vulto-estoque-importacao'))")
            cursor.execute("SELECT version_num FROM public.alembic_version")
            revisao = cursor.fetchone()
            if not revisao or revisao[0] != "20261022190000":
                raise ImportacaoEstoqueError(
                    "Aplique a migration Alembic 20261022190000 antes da carga"
                )
            cursor.execute(
                "SELECT 1 FROM public.importacao_lote WHERE codigo = %s FOR UPDATE",
                (CODIGO_LOTE,),
            )
            if cursor.fetchone():
                raise ImportacaoEstoqueError(
                    f"O lote {CODIGO_LOTE} ja existe; a carga nao sera repetida"
                )

            codigos_loja = list(CODIGO_LOJA_POR_ID_VULTO.values())
            cursor.execute(
                "SELECT codigo, id_loja FROM public.loja WHERE codigo = ANY(%s)",
                (codigos_loja,),
            )
            lojas = {codigo: id_loja for codigo, id_loja in cursor.fetchall()}
            if set(lojas) != set(codigos_loja):
                faltantes = sorted(set(codigos_loja) - set(lojas))
                raise ImportacaoEstoqueError(
                    "Lojas Vulto ausentes no Supabase: " + ", ".join(faltantes)
                )

            cursor.execute("SELECT id_unidade FROM public.unidade WHERE nome = 'vulto'")
            unidade = cursor.fetchone()
            if not unidade:
                raise ImportacaoEstoqueError("Unidade 'vulto' ausente no Supabase")
            id_unidade = unidade[0]

            ids_loja = list(lojas.values())
            cursor.execute(
                "SELECT count(*) FROM public.estoque WHERE id_loja = ANY(%s)",
                (ids_loja,),
            )
            estoque_existente = cursor.fetchone()[0]
            cursor.execute(
                "SELECT count(*) FROM public.movimentacao_estoque WHERE id_loja = ANY(%s)",
                (ids_loja,),
            )
            movimentos_existentes = cursor.fetchone()[0]
            if estoque_existente or movimentos_existentes:
                raise ImportacaoEstoqueError(
                    "Ja ha dados de estoque Vulto no Supabase "
                    f"(estoque={estoque_existente}, movimentos={movimentos_existentes}); "
                    "interrompi para nao sobrescrever nem duplicar"
                )

            cursor.execute("SELECT sku FROM public.variacao_produto")
            skus_existentes = {linha[0] for linha in cursor.fetchall()}
            cursor.execute(
                """
                SELECT codigo, id_tipo_movimentacao_estoque
                FROM public.tipo_movimentacao_estoque
                WHERE codigo = ANY(%s)
                """,
                (["entrada", "transferencia_entrada", "transferencia_saida"],),
            )
            tipos = {codigo: id_tipo for codigo, id_tipo in cursor.fetchall()}
            tipos_necessarios = {"entrada", "transferencia_entrada", "transferencia_saida"}
            if tipos_necessarios - set(tipos):
                raise ImportacaoEstoqueError(
                    "Tipos de movimento ausentes no Supabase: "
                    + ", ".join(sorted(tipos_necessarios - set(tipos)))
                )
            skus_necessarios = {
                _sku_linha(linha) for linha in fontes["estoque_conferencia"]
            } | {_sku_linha(linha) for linha in fontes["entradas_conferencia"]}
            produtos, variacoes_novas = preparar_catalogo(
                fontes["catalogo_adicoes_final"],
                fontes["itens_sem_sku_real"],
                skus_necessarios,
                skus_existentes,
            )

            produtos_por_chave = {}
            for produto in produtos:
                chave = (produto["marca"], produto["nome"])
                produtos_por_chave[chave] = _produto_variacao_no_banco(
                    cursor, produto, id_unidade
                )
            for variacao in variacoes_novas:
                produto_id = produtos_por_chave[variacao["produto_chave"]]
                cursor.execute(
                    """
                    SELECT sku FROM public.variacao_produto
                    WHERE id_produto = %s AND cor = %s AND tamanho = %s
                    """,
                    (produto_id, variacao["cor"], variacao["tamanho"]),
                )
                grade_existente = cursor.fetchone()
                if grade_existente:
                    raise ImportacaoEstoqueError(
                        "A grade aprovada colide com uma variacao existente: "
                        f"{grade_existente[0]} e {variacao['sku']}"
                    )
                cursor.execute(
                    """
                    INSERT INTO public.variacao_produto (
                        id_produto, sku, cor, tamanho, preco_venda, custo, ean, ativa
                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                    """,
                    (
                        produto_id,
                        variacao["sku"],
                        variacao["cor"],
                        variacao["tamanho"],
                        variacao["preco_venda"],
                        variacao["custo"],
                        variacao["ean"],
                        variacao["ativa"],
                    ),
                )

            cursor.execute(
                "SELECT sku, id_variacao FROM public.variacao_produto WHERE sku = ANY(%s)",
                (list(skus_necessarios),),
            )
            variacao_ids = {sku: id_variacao for sku, id_variacao in cursor.fetchall()}
            if set(variacao_ids) != skus_necessarios:
                raise ImportacaoEstoqueError("Nem todos os SKUs foram resolvidos no catalogo")

            snapshot = preparar_snapshot(fontes["estoque_conferencia"], lojas)
            snapshot = completar_snapshot_com_saldos_zero(
                snapshot, fontes["entradas_conferencia"], lojas
            )
            movimentos = preparar_movimentos(
                fontes["entradas_conferencia"], snapshot, lojas
            )
            cursor.execute(
                """
                INSERT INTO public.importacao_lote (codigo, origem)
                VALUES (%s, %s) RETURNING id_lote
                """,
                (CODIGO_LOTE, "Vulto estoque e entradas"),
            )
            id_lote = cursor.fetchone()[0]

            cursor.executemany(
                """
                INSERT INTO public.estoque (
                    id_loja, id_variacao, quantidade, estoque_minimo, atualizado_em
                )
                VALUES (%s, %s, %s, 0, %s)
                """,
                [
                    (
                        linha["id_loja"],
                        variacao_ids[linha["sku"]],
                        linha["quantidade"],
                        linha["atualizado_em"],
                    )
                    for linha in snapshot
                ],
            )

            cursor.executemany(
                """
                INSERT INTO public.movimentacao_estoque (
                    id_loja, id_variacao, id_tipo_movimentacao_estoque, quantidade,
                    quantidade_anterior, quantidade_posterior, motivo, criada_em,
                    id_lote, origem_importacao, id_origem_importacao, custo_unitario,
                    nota_fiscal, fornecedor, id_loja_destino, dados_origem
                ) VALUES (
                    %s, %s, %s, %s, %s, %s, %s, %s,
                    %s, %s, %s, %s, %s, %s, %s, %s
                )
                """,
                [
                    (
                        movimento["id_loja"],
                        variacao_ids[movimento["sku"]],
                        tipos[movimento["codigo_tipo"]],
                        movimento["quantidade"],
                        movimento["quantidade_anterior"],
                        movimento["quantidade_posterior"],
                        movimento["motivo"],
                        movimento["criada_em"],
                        id_lote,
                        ORIGEM_MOVIMENTOS,
                        movimento["id_origem"],
                        movimento["custo_unitario"],
                        movimento["nota_fiscal"],
                        movimento["fornecedor"],
                        movimento["id_loja_destino"],
                        Jsonb(
                            {
                                "linha_origem": {
                                    chave: _json_seguro(valor)
                                    for chave, valor in movimento["dados_origem"].items()
                                },
                                "saldo_reconstruido": movimento["saldo_reconstruido"],
                                "snapshot_ausente_tratado_como_zero": movimento[
                                    "snapshot_ausente_tratado_como_zero"
                                ],
                            }
                        ),
                    )
                    for movimento in movimentos
                ],
            )
    return {
        "produtos_novos": len(produtos),
        "variacoes_novas": len(variacoes_novas),
        "estoque": len(snapshot),
        "estoque_inferido_zero": sum(
            bool(linha.get("saldo_inferido_por_ausencia")) for linha in snapshot
        ),
        "movimentos": len(movimentos),
    }


def carregar_e_gravar(conexao: Any, caminho: Path) -> dict[str, int]:
    fontes = carregar_planilha(caminho)
    return gravar_importacao(conexao, fontes)
