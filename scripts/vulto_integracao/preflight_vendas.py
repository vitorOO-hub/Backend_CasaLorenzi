"""Preflight sem escrita para preparar a importacao de vendas historicas do Vulto."""

from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .clientes_vendas import indexar_compras_por_id, resolver_cpf_venda
from .lojas import CODIGO_LOJA_POR_ID_VULTO
from .valores_vendas import reconciliar_valores_venda

STATUS_PEDIDO_AUTOMATICO = {
    "aguardando": "aguardando_pagamento",
    "cancelada": "cancelado",
    "concluida": "pago",
    "estornada": "cancelado",
}
STATUS_PEDIDO_PENDENTE_DECISAO: set[str] = set()


def _linhas(aba) -> list[dict[str, Any]]:
    iterador = aba.iter_rows(values_only=True)
    cabecalhos = next(iterador, None)
    if not cabecalhos:
        return []
    return [dict(zip(cabecalhos, valores)) for valores in iterador]


def _chave(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _decimal(valor: object) -> Decimal | None:
    if valor is None or str(valor).strip() == "":
        return None
    try:
        return Decimal(str(valor).strip().replace(",", "."))
    except InvalidOperation:
        return None


def _quantidade(valor: object) -> int | None:
    numero = _decimal(valor)
    if numero is None or numero != numero.to_integral_value():
        return None
    return int(numero)


def preflight_vendas(arquivo: str | Path) -> dict[str, Any]:
    """Resume riscos e bloqueios locais; nao consulta nem altera o banco de destino."""
    workbook = load_workbook(arquivo, read_only=True, data_only=True)
    obrigatorias = {"vendas_importacao", "itens_venda_importacao", "compras_importacao"}
    faltantes = obrigatorias - set(workbook.sheetnames)
    if faltantes:
        raise ValueError("Abas ausentes: " + ", ".join(sorted(faltantes)))

    vendas = _linhas(workbook["vendas_importacao"])
    itens = _linhas(workbook["itens_venda_importacao"])
    compras = _linhas(workbook["compras_importacao"])
    cpfs_por_compra = indexar_compras_por_id(compras)
    ids_venda = {_chave(venda.get("id")) for venda in vendas}
    itens_por_venda: dict[str, list[dict[str, Any]]] = defaultdict(list)
    vendas_por_id: dict[str, dict[str, Any]] = {}
    for venda in vendas:
        vendas_por_id[_chave(venda.get("id"))] = venda
    for item in itens:
        itens_por_venda[_chave(item.get("venda_id"))].append(item)

    resolucao_cliente = Counter()
    status_legado = Counter()
    lojas_sem_mapa = Counter()
    vendas_sem_itens = 0
    for venda in vendas:
        resolucao = resolver_cpf_venda(venda, cpfs_por_compra)
        resolucao_cliente[resolucao.motivo] += 1
        status = str(venda.get("status_padronizado") or "").strip().lower()
        status_legado[status or "vazio"] += 1
        loja_id = _quantidade(venda.get("loja_id"))
        if loja_id not in CODIGO_LOJA_POR_ID_VULTO:
            lojas_sem_mapa[str(loja_id)] += 1
        if not itens_por_venda.get(_chave(venda.get("id"))):
            vendas_sem_itens += 1

    quantidade_invalida = 0
    quantidade_zero = 0
    itens_devolucao = 0
    quantidade_devolvida = 0
    preco_ausente_ou_invalido = 0
    sku_provisorio = 0
    sku_vazio = 0
    venda_item_orfao = 0
    itens_por_chave = Counter()
    por_chave_preco: dict[tuple[str, str], set[Decimal | None]] = defaultdict(set)

    for item in itens:
        venda_id = _chave(item.get("venda_id"))
        sku = str(item.get("sku_padronizado") or "").strip()
        if venda_id not in ids_venda:
            venda_item_orfao += 1
        qtd = _quantidade(item.get("qtd"))
        if qtd is None:
            quantidade_invalida += 1
        elif qtd == 0:
            quantidade_zero += 1
        elif qtd < 0:
            itens_devolucao += 1
            quantidade_devolvida += abs(qtd)
        preco = _decimal(item.get("preco_unitario_decimal"))
        if preco is None or preco < 0:
            preco_ausente_ou_invalido += 1
        if not sku:
            sku_vazio += 1
        if "provisorio" in str(item.get("pendencias") or "").lower() or sku.startswith("TMP-"):
            sku_provisorio += 1
        if venda_id and sku:
            chave = (venda_id, sku)
            itens_por_chave[chave] += 1
            por_chave_preco[chave].add(preco)

    duplicidades = [chave for chave, total in itens_por_chave.items() if total > 1]
    duplicidades_com_precos_divergentes = sum(
        len(por_chave_preco[chave]) > 1 for chave in duplicidades
    )
    canais = Counter(str(venda.get("canal") or "vazio").strip() for venda in vendas)
    lojas_canonicas = Counter(
        str(venda.get("canal_padronizado") or "vazio").strip().lower() for venda in vendas
    )
    reconciliacao_valores = Counter(
        reconciliar_valores_venda(venda, itens_por_venda.get(venda_id, []))["situacao"]
        for venda_id, venda in vendas_por_id.items()
    )

    return {
        "arquivo": Path(arquivo).name,
        "nao_gravou_no_banco": True,
        "vendas": len(vendas),
        "itens": len(itens),
        "compras_de_referencia": len(compras),
        "status_legado": dict(status_legado),
        "status_pedido_automatico": STATUS_PEDIDO_AUTOMATICO.copy(),
        "status_pendente_decisao": sorted(STATUS_PEDIDO_PENDENTE_DECISAO),
        "resolucao_cliente": dict(resolucao_cliente),
        "lojas_legado_sem_mapeamento": dict(lojas_sem_mapa),
        "vendas_sem_itens": vendas_sem_itens,
        "itens_orfaos": venda_item_orfao,
        "itens_com_quantidade_invalida": quantidade_invalida,
        "itens_com_quantidade_zero": quantidade_zero,
        "itens_identificados_como_devolucao": itens_devolucao,
        "unidades_identificadas_como_devolucao": quantidade_devolvida,
        "tratamento_devolucoes": {
            "importacao": "venda fica pendente ate definir representacao financeira da devolucao",
            "estoque": "recomposicao fica para o Bloco C; esta carga nao altera estoque",
        },
        "itens_com_preco_ausente_ou_invalido": preco_ausente_ou_invalido,
        "itens_com_sku_vazio": sku_vazio,
        "itens_com_sku_provisorio": sku_provisorio,
        "grupos_venda_sku_duplicados": len(duplicidades),
        "grupos_venda_sku_com_precos_divergentes": duplicidades_com_precos_divergentes,
        "canal_original": dict(canais),
        "loja_canonica_da_planilha": dict(lojas_canonicas),
        "reconciliacao_valores": dict(reconciliacao_valores),
        "regra_valores": {
            "desconto": "rateio proporcional nos precos dos itens somente quando fecha em centavos",
            "frete": "mantido separado em pedido.valor_frete",
            "divergencia_sem_explicacao": "pendente; nao ajustar por aproximacao",
            "devolucoes": "pendentes nesta reconciliacao financeira",
        },
        "verificacoes_pendentes_no_banco": [
            "CPF unico em public.usuario.documento",
            "SKU exato e ativo em public.variacao_produto",
            "lojas e status presentes no Supabase",
            "idempotencia por numero_pedido",
        ],
    }
