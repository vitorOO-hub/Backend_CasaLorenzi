"""Grava a planilha padronizada (uma aba por coleção, mesmas colunas do catálogo conferido) e os avisos."""
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HF = Font(name='Arial', size=10, bold=True, color='FFFFFF')
_HP = PatternFill('solid', fgColor='1F3864')
_BF = Font(name='Arial', size=10)


def _virgula(v):
    return v.replace('.', ',') if isinstance(v, str) and re.fullmatch(r'\d+\.\d{2}', v) else v


def gravar_xlsx(caminho, res):
    wb = Workbook()
    wb.remove(wb.active)
    for aba in res.abas:
        rows = [r for r in res.linhas if r['ABA'] == aba]
        com_obs = any(r['EXTRA'] for r in rows)
        cab = ['SKU', 'PRODUTO', 'CATEGORIA', 'COLEÇÃO', 'COR', 'TAM', 'PRECO', 'CUSTO', 'EAN'] + (['OBS'] if com_obs else []) + ['DESCONTINUADA']
        dados = [[r['SKU'], r['PRODUTO'], r['CATEGORIA'], r['COLECAO'], r['COR'], r['TAM'], _virgula(r['PRECO']),
                  _virgula(r['CUSTO']), r['EAN']] + ([r['EXTRA']] if com_obs else []) + [r['DESCONTINUADA']] for r in rows]
        _aba(wb.create_sheet(aba), cab, dados)
    _aba(wb.create_sheet('Avisos'), ['aba', 'linha', 'valor', 'o que aconteceu'], [list(a) for a in res.avisos])
    alertas = [[r['ABA'], r['LINHA'], r['SKU'], r['ALERTA']] for r in res.linhas if r['ALERTA']]
    _aba(wb.create_sheet('Alertas por linha'), ['aba', 'linha', 'SKU', 'alerta'], alertas)
    wb.save(caminho)


def _aba(ws, cab, dados):
    ws.append(cab)
    for d in dados:
        ws.append(d)
    for row in ws.iter_rows():
        for c in row:
            c.font = _BF
    for c in ws[1]:
        c.font, c.fill, c.alignment = _HF, _HP, Alignment(wrap_text=True, vertical='center')
    for r in range(2, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            if ws.cell(1, c).value in ('SKU', 'TAM', 'PRECO', 'CUSTO', 'EAN'):
                ws.cell(r, c).number_format = '@'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.columns, 1):
        n = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:400])
        ws.column_dimensions[get_column_letter(i)].width = min(max(9, n + 2), 70)
