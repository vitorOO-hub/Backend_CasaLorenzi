"""Gera um relatorio de divergencias de estoque Vulto sem alterar o banco."""

import argparse
import json
import re
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from vulto_integracao.conferencia_estoque import planejar_conferencia


def _chave_coluna(valor: object) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(caractere for caractere in texto if not unicodedata.combining(caractere))
    return re.sub(r"\s+", "_", texto.strip().lower())


def _linhas_aba(planilha) -> list[dict[str, Any]]:
    linhas = planilha.iter_rows(values_only=True)
    cabecalho = None
    for numero, linha in enumerate(linhas, start=1):
        chaves = [_chave_coluna(valor) for valor in linha]
        if any(chave in {"sku", "codigo", "modelo"} for chave in chaves):
            cabecalho = (chaves, numero)
            break
    if cabecalho is None:
        return []

    colunas, linha_cabecalho = cabecalho
    resultado = []
    for valores in planilha.iter_rows(min_row=linha_cabecalho + 1, values_only=True):
        if not any(valor is not None for valor in valores):
            continue
        resultado.append(
            {coluna: valor for coluna, valor in zip(colunas, valores, strict=False) if coluna}
        )
    return resultado


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("planilha_consolidada", type=Path)
    parser.add_argument("contagem_inventario", type=Path)
    parser.add_argument("--saida", type=Path, help="Salva o relatorio detalhado em JSON")
    args = parser.parse_args()

    planilha_estoque = load_workbook(args.planilha_consolidada, read_only=True, data_only=True)
    workbook = load_workbook(args.contagem_inventario, read_only=True, data_only=True)
    try:
        estoque = _linhas_aba(planilha_estoque["estoque_conferencia"])
        contagens = {
            planilha.title: _linhas_aba(planilha)
            for planilha in workbook.worksheets
        }
    finally:
        planilha_estoque.close()
        workbook.close()

    relatorio = planejar_conferencia(estoque, contagens)
    resumo = {
        "nao_gravou_no_banco": True,
        "linhas_comparadas": len(relatorio["comparacoes"]),
        "por_situacao": relatorio["resumo"],
        "estoque_sem_sku_ou_loja": len(relatorio["estoque_sem_sku_ou_loja"]),
        "contagem_sem_sku": len(relatorio["contagem_sem_sku"]),
    }
    print(json.dumps(resumo, ensure_ascii=False, indent=2))
    if args.saida:
        args.saida.parent.mkdir(parents=True, exist_ok=True)
        args.saida.write_text(
            json.dumps(relatorio, ensure_ascii=False, indent=2, default=str) + "\n",
            encoding="utf-8",
        )
        print(f"Relatorio detalhado salvo em: {args.saida}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
