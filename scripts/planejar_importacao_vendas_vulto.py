"""Gera plano de importacao Vulto em modo seco, sem conexao com o Supabase."""

import argparse
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from vulto_integracao.plano_batel import planejar_vendas_batel
from vulto_integracao.plano_importacao import planejar_importacao


def _linhas_planilha(caminho: Path, aba: str) -> list[dict[str, object]]:
    workbook = load_workbook(caminho, read_only=True, data_only=True)
    try:
        linhas = workbook[aba].iter_rows(values_only=True)
        cabecalhos = next(linhas, None)
        if not cabecalhos:
            return []
        return [dict(zip(cabecalhos, valores, strict=True)) for valores in linhas]
    finally:
        workbook.close()


def _linhas_csv(caminho: Path) -> list[dict[str, str]]:
    with caminho.open(encoding="utf-8-sig", newline="") as arquivo:
        return list(csv.DictReader(arquivo))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("planilha", type=Path)
    parser.add_argument("referencias", type=Path, help="Diretorio com os CSVs exportados")
    parser.add_argument(
        "--batel-manual",
        type=Path,
        help="Planilha opcional com vendas manuais Batel já conciliadas e mapeadas",
    )
    parser.add_argument(
        "--saida",
        type=Path,
        help="JSON detalhado do plano; pode conter dados pessoais, proteja o arquivo",
    )
    args = parser.parse_args()

    digest = hashlib.sha256()
    with args.planilha.open("rb") as arquivo:
        for bloco in iter(lambda: arquivo.read(1024 * 1024), b""):
            digest.update(bloco)
    if args.batel_manual:
        with args.batel_manual.open("rb") as arquivo:
            for bloco in iter(lambda: arquivo.read(1024 * 1024), b""):
                digest.update(bloco)
    codigo_lote = f"VULTO-VENDAS-{digest.hexdigest()[:12].upper()}"

    referencias = args.referencias
    plano = planejar_importacao(
        vendas=_linhas_planilha(args.planilha, "vendas_importacao"),
        itens=_linhas_planilha(args.planilha, "itens_venda_importacao"),
        compras=_linhas_planilha(args.planilha, "compras_importacao"),
        usuarios=_linhas_csv(referencias / "usuario_rows.csv"),
        tipos_usuario=_linhas_csv(referencias / "tipo_usuario_rows.csv"),
        lojas=_linhas_csv(referencias / "loja_rows.csv"),
        status_pedido=_linhas_csv(referencias / "status_pedido_rows.csv"),
        variacoes=_linhas_csv(referencias / "variacao_produto_rows.csv"),
        codigo_lote=codigo_lote,
    )
    if args.batel_manual:
        plano_batel = planejar_vendas_batel(
            vendas=_linhas_planilha(args.batel_manual, "vendas_padronizadas"),
            itens_mapeados=_linhas_planilha(args.batel_manual, "mapeamento_pecas"),
            lojas=_linhas_csv(referencias / "loja_rows.csv"),
            status_pedido=_linhas_csv(referencias / "status_pedido_rows.csv"),
            variacoes=_linhas_csv(referencias / "variacao_produto_rows.csv"),
            usuarios=_linhas_csv(referencias / "usuario_rows.csv"),
            tipos_usuario=_linhas_csv(referencias / "tipo_usuario_rows.csv"),
            codigo_lote=codigo_lote,
        )
        for chave in ("pedidos", "itens_pedido", "pendencias", "rastreabilidade"):
            plano[chave].extend(plano_batel[chave])
        contagens = (
            "total_vendas",
            "total_pedidos_planejados",
            "total_vendas_pendentes",
            "total_itens_planejados",
        )
        for chave in contagens:
            plano[chave] += plano_batel[chave]
        motivos = Counter(plano.get("contagem_motivos_pendencia", {}))
        motivos.update(plano_batel.get("contagem_motivos_pendencia", {}))
        plano["contagem_motivos_pendencia"] = dict(motivos)
        plano["vendas_batel_manual"] = {
            chave: valor
            for chave, valor in plano_batel.items()
            if chave not in {"pedidos", "itens_pedido", "pendencias", "rastreabilidade"}
        }
    resumo = {
        chave: valor
        for chave, valor in plano.items()
        if chave not in {"pedidos", "itens_pedido", "pendencias", "rastreabilidade"}
    }
    print(json.dumps(resumo, ensure_ascii=False, indent=2, default=str))
    if args.saida:
        args.saida.parent.mkdir(parents=True, exist_ok=True)
        conteudo = json.dumps(plano, ensure_ascii=False, indent=2, default=str) + "\n"
        args.saida.write_text(conteudo, encoding="utf-8")
        print(f"Plano detalhado salvo em: {args.saida}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
