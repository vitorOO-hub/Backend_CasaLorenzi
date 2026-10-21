"""Padroniza a planilha de clientes do Vulto+ cruzando com o banco antigo (vulto.db)."""
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime

import openpyxl

from . import banco as banco_mod
from . import pontos as pontos_mod
from . import regras as r

ABA_ANTIGOS = 'clientes antigos (NAO APAGAR)'
COLUNAS_CLIENTES = ['Nome', 'cpf', 'E-mail', 'Tel', 'Data Cadastro', 'categoria', 'Pontos', 'Obs', 'aceita_msg', 'loja_cadastro']


@dataclass
class Resultado:
    abas: dict = field(default_factory=dict)      # nome da aba -> (cabeçalho, linhas)
    avisos: list = field(default_factory=list)    # (aba, linha, problema)


def ler_aba(wb, nome):
    ws = wb[nome]
    linhas = ws.iter_rows(values_only=True)
    cabecalho = [c for c in next(linhas)]
    saida = []
    for n, valores in enumerate(linhas, start=2):
        if all(v is None or str(v).strip() == '' for v in valores):
            continue
        d = dict(zip(cabecalho, valores))
        d['_linha'] = n
        saida.append(d)
    return cabecalho, saida


def _juntar_obs(*partes):
    vistas = []
    for p in partes:
        for item in (p or '').split('; '):
            item = item.strip()
            if item and item not in vistas:
                vistas.append(item)
    return '; '.join(vistas) or None


def _primeiro(valores):
    for v in valores:
        if r.texto(v) is not None:
            return v
    return None


def corrigir_cpf(c, conhecidos):
    """CPF inválido com um dígito errado: corrige se houver exatamente um CPF válido conhecido a um dígito."""
    if r.cpf_valido(c):
        return c
    if not c or len(c) != 11:
        return None
    candidatos = [x for x in r.cpfs_a_um_digito(c) if x in conhecidos]
    return candidatos[0] if len(candidatos) == 1 else None


def limpar(caminho_planilha, caminho_banco, referencia, data_processamento=None):
    data_processamento = data_processamento or referencia.date()
    proc = data_processamento.strftime('%d/%m/%Y')
    res = Resultado()
    wb = openpyxl.load_workbook(caminho_planilha, data_only=True)
    db = banco_mod.ler(caminho_banco)

    _, clientes = ler_aba(wb, 'Clientes')
    _, resgates = ler_aba(wb, 'Resgates')
    _, antigos = ler_aba(wb, ABA_ANTIGOS)
    antigos_por_cpf = {}
    for linha in antigos:
        c = r.cpf(linha.get('CPF/CNPJ'))
        if c:
            antigos_por_cpf.setdefault(c, linha)

    # CPFs conhecidos fora da aba Clientes: servem para corrigir CPF digitado errado.
    conhecidos = set(db.cpfs)
    for linha in resgates:
        c = r.cpf(linha.get('cliente'))
        if r.cpf_valido(c):
            conhecidos.add(c)

    # 1) CPF: só números, zeros recuperados, e correção de um dígito; sem CPF válido a conta sai.
    por_cpf = {}
    cpfs_como_digitados = set()
    for linha in clientes:
        bruto = r.cpf(linha.get('cpf'))
        cpfs_como_digitados.add(bruto)
        c = corrigir_cpf(bruto, conhecidos)
        if c is None:
            res.avisos.append(('Clientes', linha['_linha'], f'CPF inválido sem correção ({bruto}): conta removida'))
            continue
        por_cpf.setdefault(c, []).append(linha)

    # 2) Uma linha por CPF: vale o primeiro valor preenchido; os pontos das contas somam.
    #    O cliente guarda os eventos de pontos e as notas; pontos, categoria e Obs saem no fim.
    contas = {}
    for c, grupo in por_cpf.items():
        col = lambda nome: [g.get(nome) for g in grupo]  # noqa: E731
        loja = r.loja(_primeiro(col('loja_cadastro')))
        tel, faltou_ddd = r.telefone(_primeiro(col('Tel')), loja)
        obs_original = r.minusculo(_primeiro(col('Obs')))
        nomes = [r.minusculo(n) for n in col('Nome') if r.texto(n)]
        nome = nomes[0] if nomes else None
        # Nome abreviado: completa com outra linha do mesmo CPF ou com a aba clientes antigos.
        if r.nome_abreviado(nome):
            antigo = r.minusculo((antigos_por_cpf.get(c) or {}).get('NOME COMPLETO'))
            for candidato in nomes[1:] + [antigo]:
                if r.completa_o_nome(nome, candidato):
                    nome = candidato
                    break
        emails = []
        for e in col('E-mail'):
            e = r.email(e)
            if e and e not in emails:
                emails.append(e)
        contas[c] = {
            'Nome': nome,
            'cpf': c,
            'E-mail': emails[0] if emails else None,
            '_emails_alternativos': emails[1:],
            'Tel': tel,
            'Data Cadastro': r.data_cadastro(_primeiro(col('Data Cadastro'))),
            'aceita_msg': r.aceita_msg(_primeiro(col('aceita_msg')), obs_original),
            'loja_cadastro': loja,
            '_faltou_ddd': faltou_ddd,
            '_obs': [r.obs(obs_original)],
            '_notas_email': [],
            '_pontos_planilha': sum(r.pontos_da_planilha(v) for v in col('Pontos')),
            '_eventos': pontos_mod.eventos_do_banco(db.movimentos.get(c, [])),
            '_compras_sem_cpf': [],
        }

    # 3) Clientes que só existem no banco: entram sem nome, com os pontos do banco.
    for c in sorted(db.movimentos):
        if c in contas or not r.cpf_valido(c):
            continue
        movs = db.movimentos[c]
        primeira = min(m.quando for m in movs).astimezone(r.SP).strftime('%d/%m/%Y')
        lojas = [db.lojas.get(m.compra.loja_id) for m in sorted(movs, key=lambda m: m.quando) if m.compra]
        contas[c] = {
            'Nome': None, 'cpf': c, 'E-mail': None, 'Tel': None, 'Data Cadastro': None,
            'aceita_msg': 'não',
            'loja_cadastro': _loja_do_banco(lojas[0] if lojas else None),
            '_faltou_ddd': False,
            '_obs': [f'cliente vindo do banco (vulto.db) em {proc}', f'1ª movimentação de pontos em {primeira}',
                     'nome não informado: pedir ao cliente'],
            '_pontos_planilha': 0,
            '_eventos': pontos_mod.eventos_do_banco(movs),
            '_compras_sem_cpf': [],
        }

    # 3b) Aba clientes antigos: quem estava ATIVO e não está na aba Clientes vem para cá com 0 pontos
    #     (os pontos de 2023 venceram). Se ele já veio do banco, os dados são mesclados.
    #     Se o CPF dele estava digitado errado na aba Clientes (e foi corrigido), os dados são mesclados.
    for c, linha in antigos_por_cpf.items():
        if (r.minusculo(linha.get('STATUS')) != 'ativo') or c in cpfs_como_digitados or not r.cpf_valido(c):
            continue
        detalhe = (f"(pontos 2023: {r.pontos_da_planilha(linha.get('PONTOS_2023'))}; "
                   f"última compra: {r.dia_br(linha.get('ULT_COMPRA'))})")
        nome = r.minusculo(linha.get('NOME COMPLETO'))
        if c in por_cpf:
            contas[c]['Nome'] = contas[c]['Nome'] or nome
            contas[c]['_obs'].insert(0, f'dados da aba clientes antigos mesclados em {proc} {detalhe}')
            continue
        if c in contas:  # já veio do banco: ganha o nome e a origem passa a ser a aba antiga
            contas[c]['Nome'] = nome
            contas[c]['_obs'] = [f'migrado da aba clientes antigos em {proc} {detalhe}']
            continue
        contas[c] = {
            'Nome': nome, 'cpf': c, 'E-mail': None, 'Tel': None, 'Data Cadastro': None,
            'aceita_msg': 'não', 'loja_cadastro': None,
            '_faltou_ddd': False,
            '_obs': [f'migrado da aba clientes antigos em {proc} {detalhe}'],
            '_pontos_planilha': 0, '_eventos': [], '_compras_sem_cpf': [],
        }

    # 3c) E-mail usado por mais de um CPF (familiares): fica com quem tem o primeiro nome igual ao
    #     início do e-mail; os outros perdem o e-mail e recebem uma nota.
    por_email_bruto = {}
    for c, conta in contas.items():
        if conta.get('E-mail'):
            por_email_bruto.setdefault(conta['E-mail'], []).append(c)
    for email_, cpfs_ in por_email_bruto.items():
        if len(cpfs_) < 2:
            continue
        dono = r.dono_do_email(email_, [contas[c]['Nome'] for c in cpfs_])
        if dono is None:
            res.avisos.append(('Clientes', '-', f'e-mail {email_} usado por {len(cpfs_)} CPFs sem dono claro'))
            continue
        for c in cpfs_:
            if contas[c]['Nome'] != dono:
                contas[c]['E-mail'] = None
                contas[c]['_notas_email'].append(
                    f'e-mail removido: {email_} é de {dono} (mesmo sobrenome, provável familiar); pedir e-mail próprio')

    # 4) Compras lançadas com telefone ou e-mail: ligadas ao cliente pelo e-mail ou pelos
    #    últimos 9 dígitos do telefone (um telefone sem DDD ganha o DDD que veio na compra).
    por_email, por_tel = {}, {}
    for c, conta in contas.items():
        if conta['E-mail']:
            por_email.setdefault(conta['E-mail'], set()).add(c)
        if conta['Tel']:
            por_tel.setdefault(conta['Tel'][-9:], set()).add(c)
    for compra in db.sem_cpf:
        if '@' in compra.cliente_ref:
            donos = por_email.get(r.email(compra.cliente_ref), set())
            tel_compra = None
        else:
            tel_compra = r.telefone_de_contato(compra.cliente_ref)
            donos = por_tel.get(tel_compra[-9:], set()) if tel_compra else set()
        if len(donos) != 1:
            continue
        conta = contas[next(iter(donos))]
        conta['_compras_sem_cpf'].append(compra)
        if tel_compra and len(tel_compra) == 11 and conta['Tel'] and len(conta['Tel']) == 9:
            conta['Tel'] = tel_compra
            conta['_faltou_ddd'] = False

    # 4b) Telefone ainda sem DDD (cadastro pelo site ou app): usa o DDD da loja física onde o
    #     cliente mais comprou, segundo o banco. Sem compra em loja física, fica sem DDD.
    ddd_da_loja_id = {i: r.DDD_DA_LOJA.get(_loja_do_banco(nome)) for i, nome in db.lojas.items()}
    for c, conta in contas.items():
        if not conta['_faltou_ddd']:
            continue
        ddds = Counter(ddd_da_loja_id.get(m.compra.loja_id) for m in db.movimentos.get(c, []) if m.compra)
        ddds.pop(None, None)
        if ddds:
            conta['Tel'] = ddds.most_common(1)[0][0] + conta['Tel']
            conta['_faltou_ddd'] = False

    # 5) Pontos, categoria e observação de cada cliente.
    saida = []
    for c, conta in contas.items():
        eventos = list(conta['_eventos'])
        creditadas, sem_pontos, vistas = [], [], set()
        for x in sorted(conta['_compras_sem_cpf'], key=lambda x: x.quando):
            # Mesma compra lançada duas vezes (com o telefone escrito de outro jeito, por exemplo):
            # mesmo valor, mesma loja e mesmo dia. Conta uma vez só.
            chave = (round(x.reais, 2), x.loja_id, x.quando.astimezone(r.SP).date())
            if not x.da_pontos or chave in vistas:
                sem_pontos.append(x)
                continue
            vistas.add(chave)
            creditadas.append(x)
        eventos += [(x.quando, x.pontos) for x in creditadas]
        notas = list(conta['_obs'])
        if conta['_faltou_ddd']:
            notas.append('telefone sem ddd: pedir ao cliente')
        if conta.get('_emails_alternativos'):
            notas.append('e-mail alternativo: ' + ', '.join(conta['_emails_alternativos']))
        notas += conta.get('_notas_email', [])
        if creditadas:
            notas.append(f'pontos de {len(creditadas)} compra(s) lançada(s) com telefone/e-mail creditados em {proc}')
        elif sem_pontos:
            notas.append(f'{len(sem_pontos)} compra(s) lançada(s) com telefone/e-mail sem pontos devidos '
                         '(cancelada, aguardando ou duplicada)')
        todos = eventos + ([(pontos_mod.CONGELAMENTO, conta['_pontos_planilha'])] if conta['_pontos_planilha'] else [])
        linha = {k: v for k, v in conta.items() if not k.startswith('_')}
        linha['categoria'] = r.categoria(pontos_mod.creditos_no_ano(eventos, referencia))
        linha['Pontos'] = pontos_mod.saldo(todos, referencia)
        linha['Obs'] = _juntar_obs(*notas)
        saida.append({col: linha.get(col) for col in COLUNAS_CLIENTES})

    res.abas['Clientes'] = (COLUNAS_CLIENTES, saida)
    res.abas[ABA_ANTIGOS] = aba_antigos(antigos, contas, proc)
    # Resgates: ainda sem tratamento (as regras dependem de decisão do time); vai como está.
    cab_resgates, linhas_resgates = ler_aba(wb, 'Resgates')
    res.abas['Resgates'] = (cab_resgates, [{k: v for k, v in x.items() if k != '_linha'} for x in linhas_resgates])
    res.abas['LEIA ISSO'] = aba_leia_isso(wb)
    res.abas['resumo categorias'] = aba_resumo(saida)
    return res


def aba_antigos(antigos, contas, proc):
    """Mesmas padronizações da aba Clientes. STATUS vazio é preenchido: ativo se o cliente está na
    aba Clientes ou tem movimentação no banco; senão inativo."""
    cabecalho = ['NOME COMPLETO', 'CPF/CNPJ', 'PONTOS_2023', 'ULT_COMPRA', 'STATUS', 'obs']
    saida = []
    for linha in antigos:
        c = r.cpf(linha.get('CPF/CNPJ'))
        status = r.minusculo(linha.get('STATUS'))
        obs = None
        if not status:
            if c in contas:
                status = 'ativo'
                obs = f'status preenchido em {proc} com base na aba clientes: cliente cadastrado e ativo hoje'
            else:
                status = 'inativo'
                obs = (f'status preenchido em {proc} com base na aba clientes: cliente não está na aba clientes '
                       'nem tem movimentação no banco desde 2025')
        saida.append({
            'NOME COMPLETO': r.minusculo(linha.get('NOME COMPLETO')),
            'CPF/CNPJ': c,
            'PONTOS_2023': r.pontos_da_planilha(linha.get('PONTOS_2023')),
            'ULT_COMPRA': r.data_cadastro(linha.get('ULT_COMPRA')),
            'STATUS': status,
            'obs': obs,
        })
    return cabecalho, saida


def aba_leia_isso(wb):
    """As regras do programa, em minúsculas como o resto da planilha."""
    ws = wb['LEIA ISSO']
    linhas = [[r.minusculo(v) for v in valores] for valores in ws.iter_rows(values_only=True)]
    cabecalho = [linhas[0][0]]
    return cabecalho, [{cabecalho[0]: x[0] if x else None} for x in linhas[1:]]


def aba_resumo(clientes):
    """Quantos clientes em cada categoria (substitui a antiga Plan1, contada à mão)."""
    total = len(clientes)
    contagem = Counter(x['categoria'] or 'sem categoria' for x in clientes)
    linhas = [{'categoria': cat, 'qtd': contagem.get(cat, 0), '% do total': contagem.get(cat, 0) / total if total else 0}
              for cat in ('bronze', 'prata', 'ouro', 'sem categoria')]
    linhas.append({'categoria': 'total', 'qtd': total, '% do total': 1 if total else 0})
    return ['categoria', 'qtd', '% do total'], linhas


def _loja_do_banco(nome):
    """'Vulto Shopping Eldorado' -> 'eldorado'; 'Vulto E-commerce' -> 'site'; 'Vulto App' -> 'app'."""
    if not nome:
        return None
    s = nome.lower().replace('vulto', '').replace('shopping', '').strip()
    return r.LOJAS.get(s, s)


def agora_sp():
    return datetime.now(r.SP)
