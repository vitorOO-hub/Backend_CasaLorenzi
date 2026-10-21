"""Grava o resultado num .xlsx com as mesmas abas da planilha final padronizada."""
import openpyxl

ORDEM = ['Clientes', 'clientes antigos (NAO APAGAR)', 'Resgates', 'LEIA ISSO', 'resumo categorias']


def gravar_xlsx(caminho, res):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nome in ORDEM + [n for n in res.abas if n not in ORDEM]:
        if nome not in res.abas:
            continue
        cabecalho, linhas = res.abas[nome]
        ws = wb.create_sheet(nome)
        ws.append(cabecalho)
        for linha in linhas:
            ws.append([linha.get(c) for c in cabecalho])
    if res.avisos:
        ws = wb.create_sheet('avisos')
        ws.append(['aba', 'linha', 'problema'])
        for aviso in res.avisos:
            ws.append(list(aviso))
    wb.save(caminho)
