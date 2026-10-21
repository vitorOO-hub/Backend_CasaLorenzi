"""Compara o resultado do clientes_vulto com a planilha final padronizada à mão.

Uso: python conferir_com_planilha_padronizada.py inicial.xlsx vulto.db final_padronizada.xlsx [AAAA-MM-DD]
A data é a de referência dos pontos (padrão: 2026-10-19, o dia em que a planilha final foi feita).
"""
import re
import sys
from collections import Counter
from datetime import datetime

import openpyxl

from clientes_vulto.limpeza import ler_aba, limpar
from clientes_vulto.regras import SP

inicial, banco, final = sys.argv[1:4]
dia = sys.argv[4] if len(sys.argv) > 4 else '2026-10-19'
referencia = datetime.fromisoformat(dia).replace(hour=23, minute=59, tzinfo=SP)

res = limpar(inicial, banco, referencia)
print('avisos:', len(res.avisos))
for a in res.avisos[:10]:
    print('  ', a)

_, esperado = ler_aba(openpyxl.load_workbook(final, data_only=True), 'Clientes')
cab, obtido = res.abas['Clientes']
esp = {str(e['cpf']).zfill(11): e for e in esperado}
obt = {o['cpf']: o for o in obtido}
print(f"\nClientes: esperado {len(esp)} | obtido {len(obt)} | só no esperado {len(esp.keys() - obt.keys())} | só no obtido {len(obt.keys() - esp.keys())}")

# Datas dentro das observações são o dia do processamento: não contam como diferença.
sem_datas = lambda s: re.sub(r'\d{2}/\d{2}/\d{4}', 'DD/MM/AAAA', s) if isinstance(s, str) else s  # noqa: E731


def igual(col, a, b):
    a = None if a in ('', None) else a
    b = None if b in ('', None) else b
    if col == 'Pontos':
        return int(a or 0) == int(b or 0)
    if col == 'Obs':
        return sem_datas(a) == sem_datas(b)
    return (str(a) if a is not None else None) == (str(b) if b is not None else None)


for col in cab:
    if col == 'cpf':
        continue
    difs = [(c, obt[c][col], esp[c][col]) for c in esp.keys() & obt.keys() if not igual(col, obt[c][col], esp[c][col])]
    total = len(esp.keys() & obt.keys())
    print(f'\n{col}: {total - len(difs)}/{total} iguais ({(total - len(difs)) / total:.1%})')
    for c, o, e in difs[:5]:
        print(f'    obtido {str(o)[:70]!r}\n    final  {str(e)[:70]!r}')
    if col == 'Obs':
        print('    padrões que faltam:', Counter(re.sub(r'\d', '9', str(e))[:60] for _, o, e in difs).most_common(6))

# Outras abas
wb_final = openpyxl.load_workbook(final, data_only=True)
cab, obtido = res.abas['clientes antigos (NAO APAGAR)']
_, esperado = ler_aba(wb_final, 'clientes antigos (NAO APAGAR)')
esp = {str(e['CPF/CNPJ']).zfill(11): e for e in esperado}
for col in cab:
    total = len(obtido)
    ok = sum(igual(col if col != 'obs' else 'Obs', o[col], esp.get(o['CPF/CNPJ'], {}).get(col)) for o in obtido)
    print(f'\nclientes antigos / {col}: {ok}/{total} iguais')

# Resgates: a final é a inicial menos os resgates removidos, na mesma ordem.
cab, obtido = res.abas['Resgates']
_, esperado = ler_aba(wb_final, 'Resgates')
chave = lambda x: (x['data'][:10], str(x['cliente']).zfill(11), x['pontos usados'])  # noqa: E731
so_obtido = [chave(o) for o in obtido if chave(o) not in {chave(e) for e in esperado}]
so_final = [chave(e) for e in esperado if chave(e) not in {chave(o) for o in obtido}]
print(f'\nResgates: esperado {len(esperado)} | obtido {len(obtido)} | só no obtido {so_obtido} | só na final {so_final}')
obtido = [o for o in obtido if chave(o) not in so_obtido]
esperado = [e for e in esperado if chave(e) not in so_final]
for col in cab:
    difs = [(o[col], e[col]) for o, e in zip(obtido, esperado) if not igual(col if col != 'observacao' else 'Obs', o[col], e[col])]
    print(f'Resgates / {col}: {len(obtido) - len(difs)}/{len(obtido)} iguais')
    for o, e in difs[:6]:
        print(f'    obtido {str(o)[:150]!r}\n    final  {str(e)[:150]!r}')

_, obtido = res.abas['LEIA ISSO']
ws = wb_final['LEIA ISSO']
esperado = [v[0] for v in ws.iter_rows(min_row=2, values_only=True) if v and v[0]]
obtido = [o[next(iter(o))] for o in obtido if o[next(iter(o))]]
print(f"\nLEIA ISSO: {sum(o == e for o, e in zip(obtido, esperado))}/{len(esperado)} linhas iguais")

_, obtido = res.abas['resumo categorias']
_, esperado = ler_aba(wb_final, 'resumo categorias')
print('\nresumo categorias (obtido x final):')
for o, e in zip(obtido, esperado):
    print(f"   {o['categoria']:14} {o['qtd']:5} {o['% do total']:.4f}  |  {e['categoria']:14} {e['qtd']:5} {float(e['% do total']):.4f}")
