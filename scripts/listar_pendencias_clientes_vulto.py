"""Exporta vendas sem cliente conciliado, sem incluir CPF no arquivo de saida."""

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook

from vulto_integracao.clientes_vendas import (
    indexar_compras_por_id,
    indexar_clientes_por_cpf,
    listar_vendas_pendentes_cliente,
)


def _linhas_planilha(caminho: Path, aba: str) -> list[dict[str, object]]:
    workbook = load_workbook(caminho, read_only=True, data_only=True)
    try:
        planilha = workbook[aba]
        linhas = planilha.iter_rows(values_only=True)
        cabecalhos = next(linhas, None)
        if not cabecalhos:
            return []
        return [dict(zip(cabecalhos, linha)) for linha in linhas]
    finally:
        workbook.close()


def _ler_csv(caminho: Path) -> list[dict[str, str]]:
    with caminho.open(encoding="utf-8-sig", newline="") as arquivo:
        return list(csv.DictReader(arquivo))


def gerar_lista_pendencias(
    planilha: Path, usuarios_csv: Path, tipos_csv: Path
) -> list[dict[str, str]]:
    vendas = _linhas_planilha(planilha, "vendas_importacao")
    compras = _linhas_planilha(planilha, "compras_importacao")
    usuarios = _ler_csv(usuarios_csv)
    tipos = _ler_csv(tipos_csv)
    tipos_por_id = {
        linha.get("id_tipo_usuario", "").strip(): linha.get("codigo", "").strip()
        for linha in tipos
    }
    return listar_vendas_pendentes_cliente(
        vendas,
        indexar_compras_por_id(compras),
        indexar_clientes_por_cpf(usuarios, tipos_por_id),
    )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("planilha", type=Path)
    parser.add_argument("usuarios_csv", type=Path)
    parser.add_argument("tipos_csv", type=Path)
    parser.add_argument("saida", type=Path)
    args = parser.parse_args()

    pendencias = gerar_lista_pendencias(args.planilha, args.usuarios_csv, args.tipos_csv)
    args.saida.parent.mkdir(parents=True, exist_ok=True)
    with args.saida.open("w", encoding="utf-8-sig", newline="") as arquivo:
        writer = csv.DictWriter(
            arquivo, fieldnames=("venda_id", "pendencia_cliente")
        )
        writer.writeheader()
        writer.writerows(pendencias)
    print(f"Vendas mantidas pendentes: {len(pendencias)}")
    print("A saida contem somente venda_id e motivo; nenhum CPF foi exportado.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
